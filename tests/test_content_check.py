import docx
import pymupdf

from ingest.content_check import MIN_TEXT_CHARS, check_file

LINE = "Grading: homework 30 percent, midterm 30 percent, final exam 40 percent."


def make_pdf(path, lines):
    doc = pymupdf.open()
    page = doc.new_page()
    for i, line in enumerate(lines):
        page.insert_text((72, 72 + 14 * i), line)  # insert_text doesn't wrap, so one call per line
    doc.save(path)
    doc.close()


def make_docx(path, paragraphs=(), table_cells=()):
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table_cells:
        t = d.add_table(rows=1, cols=len(table_cells))
        for i, text in enumerate(table_cells):
            t.cell(0, i).text = text
    d.save(path)


def test_pdf_with_text_has_text(tmp_path):
    p = tmp_path / "s.pdf"
    make_pdf(p, [LINE] * 10)
    r = check_file(p, "pdf")
    assert r.page_count == 1
    assert r.text_chars >= MIN_TEXT_CHARS
    assert r.has_text is True
    assert r.error is None


def test_blank_pdf_has_no_text(tmp_path):
    # Stand-in for a scanned PDF: pages exist, but there is no extractable text.
    p = tmp_path / "blank.pdf"
    make_pdf(p, [])
    r = check_file(p, "pdf")
    assert r.page_count == 1
    assert r.text_chars == 0
    assert r.has_text is False


def test_docx_without_tables_still_counts_paragraphs(tmp_path):
    # Regression test: `chars` was once assigned inside the table loop, so a
    # DOCX with no tables raised UnboundLocalError.
    p = tmp_path / "plain.docx"
    make_docx(p, paragraphs=[LINE] * 5)
    r = check_file(p, "docx")
    assert r.error is None
    assert r.has_text is True
    assert r.page_count is None  # DOCX has no fixed pages


def test_docx_counts_table_cells(tmp_path):
    without = tmp_path / "a.docx"
    with_table = tmp_path / "b.docx"
    make_docx(without, paragraphs=[LINE] * 5)
    make_docx(with_table, paragraphs=[LINE] * 5, table_cells=["Midterm", "30%"])
    diff = check_file(with_table, "docx").text_chars - check_file(without, "docx").text_chars
    assert diff == len("Midterm30%")  # grading tables must be counted


def test_corrupt_file_reports_error_instead_of_crashing(tmp_path):
    p = tmp_path / "fake.pdf"
    p.write_text("not a pdf")
    r = check_file(p, "pdf")
    # None, not False: "couldn't check" is different from "checked, no text".
    assert r.has_text is None
    assert r.error


def test_unknown_extension_is_unknown(tmp_path):
    p = tmp_path / "old.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0")
    r = check_file(p, "doc")
    assert (r.page_count, r.text_chars, r.has_text, r.error) == (None, None, None, None)
