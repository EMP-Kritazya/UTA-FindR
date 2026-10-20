"""Lightweight "does this file have text?" check, run right after download.

Not parsing: Docling does that later. This only flags files that are empty or
scanned (images only), which would need OCR before they're usable.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

MIN_TEXT_CHARS = 200  # We need to extract at least 200 characters for it to be a valid document (can be tuned)


@dataclass
class ContentCheck:
    page_count: int | None
    text_chars: int | None
    has_text: bool | None
    error: str | None = None


def _count(text: str) -> int:
    """Count non-whitespace characters."""
    return sum(1 for c in text if not c.isspace())
        

def check_file(path: Path, ext: str) -> ContentCheck:
    try:
        if ext == "pdf":
            # PDF
            import pymupdf

            with pymupdf.open(path) as doc:
                total = sum(_count(page.get_text()) for page in doc)
                return ContentCheck(doc.page_count, total, total>=MIN_TEXT_CHARS)

        if ext == "docx":
                # DOCX
                import docx
                d = docx.Document(str(path))

                parts = [p.text for p in d.paragraphs]

                for table in d.tables:
                    for row in table.rows:
                        parts.extend(cell.text for cell in row.cells)
                        
                chars = _count("\n".join(parts))
                
                return ContentCheck(None, chars, chars >= MIN_TEXT_CHARS)
        return ContentCheck(None, None, None)
    except Exception as e:
        return ContentCheck(None, None, None, error=f"{type(e).__name__}: {e}")